#!/usr/bin/env python3
from __future__ import annotations

import csv
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[3]
INPUT_XLSX = REPO_ROOT / "apps" / "data" / "raw" / "industry" / "abs_lfs_table05_feb2026.xlsx"
OUTPUT_CSV = REPO_ROOT / "apps" / "data" / "raw" / "industry" / "abs_industry_top5_state_latest.csv"

GEO_TO_CODE = {
    "Australia": "AU",
    "New South Wales": "NSW",
    "Victoria": "VIC",
    "Queensland": "QLD",
    "South Australia": "SA",
    "Western Australia": "WA",
    "Tasmania": "TAS",
    "Northern Territory": "NT",
    "Australian Capital Territory": "ACT",
}

DESCRIPTION_RE = re.compile(r"^(?P<geo>[^;]+?)\s*;\s*(?P<industry>[^;]+?)\s*;\s*Employed total\s*;$")


def normalize_geo(raw_geo: str) -> str:
    return raw_geo.strip().lstrip(">").strip()


def build_industry_series(workbook_path: Path) -> list[dict]:
    workbook = load_workbook(workbook_path, data_only=True)
    rows: list[dict] = []
    for sheet_name in ("Data1", "Data2", "Data3"):
        worksheet = workbook[sheet_name]
        for column in range(2, worksheet.max_column + 1):
            descriptor = worksheet.cell(1, column).value
            series_type = worksheet.cell(3, column).value
            series_id = worksheet.cell(10, column).value
            if not descriptor or not series_id:
                continue
            match = DESCRIPTION_RE.match(str(descriptor).strip())
            if not match:
                continue
            if str(series_type).strip() != "Original":
                continue

            geo_name = normalize_geo(match.group("geo"))
            industry_name = match.group("industry").strip()
            values: dict[datetime, float] = {}
            for row in range(11, worksheet.max_row + 1):
                date_cell = worksheet.cell(row, 1).value
                value_cell = worksheet.cell(row, column).value
                if isinstance(date_cell, datetime) and value_cell is not None:
                    values[date_cell] = float(value_cell)
            if not values:
                continue
            latest_date = max(values)
            rows.append(
                {
                    "geo_name": geo_name,
                    "industry_name": industry_name,
                    "reference_date": latest_date,
                    "employed_thousands": values[latest_date],
                    "series_id": str(series_id),
                    "series_type": str(series_type),
                }
            )
    return rows


def main() -> None:
    if not INPUT_XLSX.exists():
        raise FileNotFoundError(f"Missing source workbook: {INPUT_XLSX}")

    series_rows = build_industry_series(INPUT_XLSX)
    by_geo: dict[str, list[dict]] = {}
    for row in series_rows:
        geo_name = row["geo_name"]
        code = GEO_TO_CODE.get(geo_name)
        if not code:
            continue
        by_geo.setdefault(code, []).append(row)

    output_rows: list[dict[str, str]] = []
    for geo_code, rows in by_geo.items():
        if not rows:
            continue
        latest_date = max(row["reference_date"] for row in rows)
        latest_rows = [row for row in rows if row["reference_date"] == latest_date]
        total_employed = sum(row["employed_thousands"] for row in latest_rows)
        if total_employed <= 0:
            continue
        ranked = sorted(latest_rows, key=lambda item: item["employed_thousands"], reverse=True)[:5]
        for rank, row in enumerate(ranked, start=1):
            share = (row["employed_thousands"] / total_employed) * 100.0
            output_rows.append(
                {
                    "geo_code": geo_code,
                    "reference_month": latest_date.strftime("%Y-%m"),
                    "industry_rank": str(rank),
                    "industry_name": row["industry_name"],
                    "employment_share": f"{share:.1f}",
                    "employed_persons": str(int(round(row["employed_thousands"] * 1000))),
                    "total_employed_persons": str(int(round(total_employed * 1000))),
                    "series_id": row["series_id"],
                    "series_type": row["series_type"],
                }
            )

    fieldnames = [
        "geo_code",
        "reference_month",
        "industry_rank",
        "industry_name",
        "employment_share",
        "employed_persons",
        "total_employed_persons",
        "series_id",
        "series_type",
    ]
    OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT_CSV.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(output_rows)
    print(f"Wrote {OUTPUT_CSV} with {len(output_rows)} rows.")


if __name__ == "__main__":
    main()
