#!/usr/bin/env python3
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[3]
INPUT_XLSX = REPO_ROOT / "apps" / "data" / "raw" / "labour_market" / "abs_lfs_table12_mar2026.xlsx"
OUTPUT_CSV = REPO_ROOT / "apps" / "data" / "raw" / "labour_market" / "abs_lfs_state_latest.csv"

GEOGRAPHIES = {
    "AU": "Australia",
    "NSW": "> New South Wales",
    "VIC": "> Victoria",
    "QLD": "> Queensland",
    "SA": "> South Australia",
    "WA": "> Western Australia",
    "TAS": "> Tasmania",
    "NT": "> Northern Territory",
    "ACT": "> Australian Capital Territory",
}


def build_series_index(workbook_path: Path) -> dict[str, dict]:
    workbook = load_workbook(workbook_path, data_only=True)
    series_by_id: dict[str, dict] = {}
    for sheet_name in ("Data1", "Data2", "Data3"):
        worksheet = workbook[sheet_name]
        for column in range(2, worksheet.max_column + 1):
            descriptor = worksheet.cell(1, column).value
            unit = worksheet.cell(2, column).value
            series_type = worksheet.cell(3, column).value
            series_id = worksheet.cell(10, column).value
            if not descriptor or not series_id:
                continue

            values: dict[datetime, float] = {}
            for row in range(11, worksheet.max_row + 1):
                date_cell = worksheet.cell(row, 1).value
                value_cell = worksheet.cell(row, column).value
                if isinstance(date_cell, datetime) and value_cell is not None:
                    values[date_cell] = float(value_cell)

            if values:
                series_by_id[str(series_id)] = {
                    "descriptor": str(descriptor),
                    "unit": str(unit),
                    "series_type": str(series_type),
                    "values": values,
                }
    return series_by_id


def select_series(
    series_catalog: dict[str, dict],
    descriptor: str,
    preferences: tuple[str, ...],
) -> tuple[str, dict]:
    matching = [(series_id, entry) for series_id, entry in series_catalog.items() if entry["descriptor"] == descriptor]
    if not matching:
        raise KeyError(f"Series not found for descriptor: {descriptor}")
    by_type = {entry["series_type"]: (series_id, entry) for series_id, entry in matching}
    for preferred in preferences:
        if preferred in by_type:
            return by_type[preferred]
    return matching[0]


def extract_rows(series_catalog: dict[str, dict]) -> list[dict]:
    rows: list[dict] = []
    for code, area in GEOGRAPHIES.items():
        employment_id, employment = select_series(
            series_catalog,
            f"Employed total ;  Persons ;  {area} ;",
            preferences=("Seasonally Adjusted", "Trend", "Original"),
        )
        unemployment_id, unemployment = select_series(
            series_catalog,
            f"Unemployment rate ;  Persons ;  {area} ;",
            preferences=("Seasonally Adjusted", "Trend", "Original"),
        )
        population_id, population = select_series(
            series_catalog,
            f"Civilian population aged 15 years and over ;  Persons ;  {area} ;",
            preferences=("Original", "Trend", "Seasonally Adjusted"),
        )

        common_dates = set(employment["values"]) & set(unemployment["values"]) & set(population["values"])
        if not common_dates:
            raise ValueError(f"No shared date for geography {code}")
        latest_date = max(common_dates)

        rows.append(
            {
                "geo_code": code,
                "reference_month": latest_date.strftime("%Y-%m"),
                "population_15_plus": str(int(round(population["values"][latest_date] * 1000))),
                "employment": str(int(round(employment["values"][latest_date] * 1000))),
                "unemployment_rate": f"{unemployment['values'][latest_date]:.1f}",
                "employment_series_type": employment["series_type"],
                "unemployment_rate_series_type": unemployment["series_type"],
                "population_series_type": population["series_type"],
                "employment_series_id": employment_id,
                "unemployment_rate_series_id": unemployment_id,
                "population_series_id": population_id,
            }
        )
    return rows


def main() -> None:
    if not INPUT_XLSX.exists():
        raise FileNotFoundError(f"Missing source workbook: {INPUT_XLSX}")

    catalog = build_series_index(INPUT_XLSX)
    rows = extract_rows(catalog)
    OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "geo_code",
        "reference_month",
        "population_15_plus",
        "employment",
        "unemployment_rate",
        "employment_series_type",
        "unemployment_rate_series_type",
        "population_series_type",
        "employment_series_id",
        "unemployment_rate_series_id",
        "population_series_id",
    ]
    with OUTPUT_CSV.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    print(f"Wrote {OUTPUT_CSV} with {len(rows)} rows.")


if __name__ == "__main__":
    main()
