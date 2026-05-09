from __future__ import annotations

import csv
import json
import re
from enum import Enum
from functools import lru_cache
from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field, HttpUrl, ValidationError, field_validator

from .schemas import (
    DataStatus,
    IndustryStat,
    LabourMarketSnapshot,
    Occupation,
    ShortageCategory,
    SourceReference,
)

REPO_ROOT = Path(__file__).resolve().parents[3]
DATA_ROOT = REPO_ROOT / "apps" / "data"
SOURCE_METADATA_FILE = DATA_ROOT / "metadata" / "sources.json"

REAL_RATING_TO_CATEGORY = {
    "NS": ShortageCategory.NO_SHORTAGE,
    "S": ShortageCategory.SHORTAGE,
    "R": ShortageCategory.REGIONAL_SHORTAGE,
    "M": ShortageCategory.METROPOLITAN_SHORTAGE,
}

STATE_COLUMNS = {
    "nsw": "NSW",
    "vic": "VIC",
    "qld": "QLD",
    "sa": "SA",
    "wa": "WA",
    "tas": "TAS",
    "nt": "NT",
    "act": "ACT",
}

MAJOR_GROUP_LABELS = {
    1: "Managers",
    2: "Professionals",
    3: "Technicians and Trades Workers",
    4: "Community and Personal Service Workers",
    5: "Clerical and Administrative Workers",
    6: "Sales Workers",
    7: "Machinery Operators and Drivers",
    8: "Labourers",
    9: "Other",
}

LEGACY_OCCUPATION_IDS_BY_CODE = {
    "261313": "software-engineer",
    "261111": "ict-business-analyst",
    "233211": "civil-engineer",
    "254499": "registered-nurse",
    "331212": "carpenter",
}


class FileFormat(str, Enum):
    CSV = "csv"
    JSON = "json"
    XLSX = "xlsx"


class SourceMetadata(BaseModel):
    dataset_id: str
    dataset_name: str
    provider: str
    source_url: HttpUrl
    download_url: HttpUrl | None = None
    local_path: str
    file_format: FileFormat
    last_updated: str
    data_status: DataStatus = DataStatus.REAL
    description: str = ""
    license: str | None = None
    sheet_name: str | None = None
    header_row: int = Field(default=1, ge=1)


class OccupationShortageRow(BaseModel):
    anzsco_code: str = Field(pattern=r"^\d{6}$")
    occupation_title: str = Field(min_length=2)
    national_rating: str
    nsw: str
    vic: str
    qld: str
    sa: str
    wa: str
    tas: str
    nt: str
    act: str
    skill_level: int = Field(ge=1, le=5)
    major_group: int = Field(ge=1, le=9)

    @field_validator("anzsco_code", mode="before")
    @classmethod
    def normalize_code(cls, value: Any) -> str:
        return str(value).strip()

    @field_validator(
        "national_rating",
        "nsw",
        "vic",
        "qld",
        "sa",
        "wa",
        "tas",
        "nt",
        "act",
        mode="before",
    )
    @classmethod
    def normalize_rating(cls, value: Any) -> str:
        rating = str(value).strip().upper()
        if rating not in REAL_RATING_TO_CATEGORY:
            raise ValueError(f"invalid shortage rating: {rating}")
        return rating

    @field_validator("skill_level", "major_group", mode="before")
    @classmethod
    def parse_int_fields(cls, value: Any) -> int:
        return int(str(value).strip())


class RealOccupationDataset(BaseModel):
    source: SourceReference
    occupations: list[Occupation]
    shortage_matrix: dict[str, dict[str, ShortageCategory]]


class LabourMarketRow(BaseModel):
    geo_code: str = Field(pattern=r"^(AU|NSW|VIC|QLD|SA|WA|TAS|NT|ACT)$")
    reference_month: str = Field(pattern=r"^\d{4}-\d{2}$")
    population_15_plus: int = Field(ge=0)
    employment: int = Field(ge=0)
    unemployment_rate: float = Field(ge=0, le=100)

    @field_validator("geo_code", mode="before")
    @classmethod
    def normalize_geo_code(cls, value: Any) -> str:
        return str(value).strip().upper()

    @field_validator("population_15_plus", "employment", mode="before")
    @classmethod
    def parse_population_like_int(cls, value: Any) -> int:
        return int(float(str(value).strip()))

    @field_validator("unemployment_rate", mode="before")
    @classmethod
    def parse_unemployment_rate(cls, value: Any) -> float:
        return float(str(value).strip())


class RealLabourMarketDataset(BaseModel):
    source: SourceReference
    latest_month: str
    by_geo_code: dict[str, LabourMarketSnapshot]


class IndustryRow(BaseModel):
    geo_code: str = Field(pattern=r"^(AU|NSW|VIC|QLD|SA|WA|TAS|NT|ACT)$")
    reference_month: str = Field(pattern=r"^\d{4}-\d{2}$")
    industry_rank: int = Field(ge=1, le=20)
    industry_name: str = Field(min_length=2)
    employment_share: float = Field(ge=0, le=100)

    @field_validator("geo_code", mode="before")
    @classmethod
    def normalize_industry_geo_code(cls, value: Any) -> str:
        return str(value).strip().upper()

    @field_validator("industry_rank", mode="before")
    @classmethod
    def parse_rank(cls, value: Any) -> int:
        return int(str(value).strip())

    @field_validator("employment_share", mode="before")
    @classmethod
    def parse_share(cls, value: Any) -> float:
        return float(str(value).strip())


class RealIndustryDataset(BaseModel):
    source: SourceReference
    latest_month: str
    by_geo_code: dict[str, list[IndustryStat]]


class DataLoadError(ValueError):
    pass


def _slugify(value: str) -> str:
    normalized = value.lower().replace("&", " and ").replace("/", " ")
    slug = re.sub(r"[^a-z0-9]+", "-", normalized).strip("-")
    return re.sub(r"-{2,}", "-", slug)


def _row_values(row: list[Any]) -> list[str]:
    cleaned: list[str] = []
    for cell in row:
        if cell is None:
            cleaned.append("")
            continue
        if isinstance(cell, str):
            cleaned.append(cell.strip())
            continue
        cleaned.append(str(cell).strip())
    return cleaned


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        return [{key: (value or "").strip() for key, value in row.items()} for row in reader]


def _read_json(path: Path) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, list):
        return payload
    if isinstance(payload, dict):
        rows = payload.get("rows")
        if isinstance(rows, list):
            return rows
    raise DataLoadError(f"Unsupported JSON table format in {path}")


def _read_xlsx(path: Path, sheet_name: str | None, header_row: int) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ModuleNotFoundError as error:  # pragma: no cover - environment dependent
        raise DataLoadError(
            "XLSX source requires openpyxl. Install it with: python3 -m pip install --user openpyxl"
        ) from error

    workbook = load_workbook(path, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    header_cells = _row_values(list(worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))[0])
    if not any(header_cells):
        raise DataLoadError(f"Empty header row {header_row} in {path}")

    rows: list[dict[str, str]] = []
    for raw in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        values = _row_values(list(raw))
        if not any(values):
            continue
        item = {
            header_cells[index]: values[index]
            for index in range(min(len(header_cells), len(values)))
            if header_cells[index]
        }
        if item:
            rows.append(item)
    return rows


def read_table(path: Path, file_format: FileFormat, *, sheet_name: str | None = None, header_row: int = 1) -> list[dict[str, Any]]:
    if file_format == FileFormat.CSV:
        return _read_csv(path)
    if file_format == FileFormat.JSON:
        return _read_json(path)
    if file_format == FileFormat.XLSX:
        return _read_xlsx(path, sheet_name=sheet_name, header_row=header_row)
    raise DataLoadError(f"Unsupported file format: {file_format}")


def _load_source_metadata(dataset_id: str) -> SourceMetadata:
    payload = json.loads(SOURCE_METADATA_FILE.read_text(encoding="utf-8"))
    for item in payload.get("sources", []):
        if item.get("dataset_id") == dataset_id:
            return SourceMetadata.model_validate(item)
    raise DataLoadError(f"Dataset id not found in source metadata: {dataset_id}")


def _build_occupation_id(code: str, title: str, used_ids: set[str]) -> str:
    legacy = LEGACY_OCCUPATION_IDS_BY_CODE.get(code)
    if legacy:
        used_ids.add(legacy)
        return legacy

    base = f"{_slugify(title)}-{code}"
    candidate = base
    serial = 2
    while candidate in used_ids:
        candidate = f"{base}-{serial}"
        serial += 1
    used_ids.add(candidate)
    return candidate


@lru_cache(maxsize=1)
def load_real_occupation_dataset(dataset_id: str = "jsa_osl_2025_anzsco2022") -> RealOccupationDataset:
    metadata = _load_source_metadata(dataset_id)
    local_path = REPO_ROOT / metadata.local_path
    if not local_path.exists():
        raise DataLoadError(f"Source file missing: {local_path}")

    raw_rows = read_table(
        local_path,
        metadata.file_format,
        sheet_name=metadata.sheet_name,
        header_row=metadata.header_row,
    )
    if not raw_rows:
        raise DataLoadError(f"Source file has no rows: {local_path}")

    occupations: list[Occupation] = []
    shortage_matrix: dict[str, dict[str, ShortageCategory]] = {}
    used_ids: set[str] = set()
    errors: list[str] = []

    for index, row in enumerate(raw_rows, start=2):
        try:
            parsed = OccupationShortageRow.model_validate(row)
        except ValidationError as error:
            errors.append(f"row {index}: {error.errors()}")
            if len(errors) >= 10:
                break
            continue

        occupation_id = _build_occupation_id(parsed.anzsco_code, parsed.occupation_title, used_ids)
        occupation = Occupation(
            id=occupation_id,
            label=parsed.occupation_title,
            anzsco_code=parsed.anzsco_code,
            major_group=MAJOR_GROUP_LABELS.get(parsed.major_group, f"Major Group {parsed.major_group}"),
            skill_level=parsed.skill_level,
        )
        occupations.append(occupation)

        shortage_matrix[occupation_id] = {
            "AU": REAL_RATING_TO_CATEGORY[parsed.national_rating],
            **{
                state_code: REAL_RATING_TO_CATEGORY[getattr(parsed, column)]
                for column, state_code in STATE_COLUMNS.items()
            },
        }

    if errors:
        joined = "\n".join(errors)
        raise DataLoadError(f"Invalid rows detected in {local_path}:\n{joined}")

    source_reference = SourceReference(
        title=f"{metadata.provider} - {metadata.dataset_name}",
        url=str(metadata.source_url),
        last_updated=metadata.last_updated,
        data_status=metadata.data_status,
    )
    return RealOccupationDataset(
        source=source_reference,
        occupations=occupations,
        shortage_matrix=shortage_matrix,
    )


@lru_cache(maxsize=1)
def load_real_labour_market_dataset(
    dataset_id: str = "abs_lfs_table12_state_latest",
) -> RealLabourMarketDataset:
    metadata = _load_source_metadata(dataset_id)
    local_path = REPO_ROOT / metadata.local_path
    if not local_path.exists():
        raise DataLoadError(f"Source file missing: {local_path}")

    raw_rows = read_table(
        local_path,
        metadata.file_format,
        sheet_name=metadata.sheet_name,
        header_row=metadata.header_row,
    )
    if not raw_rows:
        raise DataLoadError(f"Source file has no rows: {local_path}")

    by_geo_code: dict[str, LabourMarketSnapshot] = {}
    latest_month = ""
    errors: list[str] = []

    for index, row in enumerate(raw_rows, start=2):
        try:
            parsed = LabourMarketRow.model_validate(row)
        except ValidationError as error:
            errors.append(f"row {index}: {error.errors()}")
            if len(errors) >= 10:
                break
            continue

        latest_month = max(latest_month, parsed.reference_month)
        by_geo_code[parsed.geo_code] = LabourMarketSnapshot(
            population=parsed.population_15_plus,
            employment=parsed.employment,
            unemployment_rate=parsed.unemployment_rate,
        )

    if errors:
        joined = "\n".join(errors)
        raise DataLoadError(f"Invalid rows detected in {local_path}:\n{joined}")

    if not by_geo_code:
        raise DataLoadError(f"No valid labour market rows found in {local_path}")

    source_reference = SourceReference(
        title=f"{metadata.provider} - {metadata.dataset_name}",
        url=str(metadata.source_url),
        last_updated=metadata.last_updated,
        data_status=metadata.data_status,
    )
    return RealLabourMarketDataset(
        source=source_reference,
        latest_month=latest_month,
        by_geo_code=by_geo_code,
    )


@lru_cache(maxsize=1)
def load_real_industry_dataset(
    dataset_id: str = "abs_lfs_table05_industry_top5_state_latest",
) -> RealIndustryDataset:
    metadata = _load_source_metadata(dataset_id)
    local_path = REPO_ROOT / metadata.local_path
    if not local_path.exists():
        raise DataLoadError(f"Source file missing: {local_path}")

    raw_rows = read_table(
        local_path,
        metadata.file_format,
        sheet_name=metadata.sheet_name,
        header_row=metadata.header_row,
    )
    if not raw_rows:
        raise DataLoadError(f"Source file has no rows: {local_path}")

    latest_month = ""
    staged: dict[str, list[tuple[int, IndustryStat]]] = {}
    errors: list[str] = []

    for index, row in enumerate(raw_rows, start=2):
        try:
            parsed = IndustryRow.model_validate(row)
        except ValidationError as error:
            errors.append(f"row {index}: {error.errors()}")
            if len(errors) >= 10:
                break
            continue

        latest_month = max(latest_month, parsed.reference_month)
        staged.setdefault(parsed.geo_code, []).append(
            (
                parsed.industry_rank,
                IndustryStat(
                    name=parsed.industry_name,
                    employment_share=parsed.employment_share,
                ),
            )
        )

    if errors:
        joined = "\n".join(errors)
        raise DataLoadError(f"Invalid rows detected in {local_path}:\n{joined}")

    if not staged:
        raise DataLoadError(f"No valid industry rows found in {local_path}")

    by_geo_code = {
        geo_code: [stat for _, stat in sorted(rows, key=lambda pair: pair[0])]
        for geo_code, rows in staged.items()
    }

    source_reference = SourceReference(
        title=f"{metadata.provider} - {metadata.dataset_name}",
        url=str(metadata.source_url),
        last_updated=metadata.last_updated,
        data_status=metadata.data_status,
    )
    return RealIndustryDataset(
        source=source_reference,
        latest_month=latest_month,
        by_geo_code=by_geo_code,
    )
